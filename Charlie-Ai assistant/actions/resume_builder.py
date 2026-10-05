"""Create a clean DOCX resume from remembered context plus a short brief."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from core.undo import push_undo
from memory.memory_manager import load_memory


def _value(entry) -> str:
    return str(entry.get("value", "") if isinstance(entry, dict) else entry or "").strip()


def _remove(path: Path) -> str:
    if path.exists():
        path.unlink()
        return f"Removed {path.name}."
    return f"{path.name} was already missing."


def _remembered_defaults() -> dict:
    memory = load_memory()
    identity = memory.get("identity", {}) or {}
    projects = memory.get("projects", {}) or {}
    return {
        "full_name": _value(identity.get("name")),
        "headline": _value(identity.get("job")),
        "city": _value(identity.get("city")),
        "projects": [f"{key.replace('_', ' ').title()}: {_value(entry)}"
                     for key, entry in list(projects.items())[:6] if _value(entry)],
    }


def resume_builder(parameters: dict, player=None, **_unused) -> str:
    params = parameters or {}
    defaults = _remembered_defaults()
    name = str(params.get("full_name") or defaults["full_name"]).strip()
    target_role = str(params.get("target_role") or defaults["headline"]).strip()
    if not name:
        return "I only need your full name before I can create the resume."
    if not target_role:
        return "I only need the job role you want this resume to target."

    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt, RGBColor

        out_raw = str(params.get("output_path") or "").strip().strip('"')
        if out_raw:
            output = Path(out_raw).expanduser()
            if output.suffix.lower() != ".docx":
                output = output.with_suffix(".docx")
        else:
            output = Path.home() / "Documents" / f"{name.replace(' ', '_')}_Resume.docx"
        output.parent.mkdir(parents=True, exist_ok=True)
        if output.exists():
            output = output.with_name(f"{output.stem}_{datetime.now():%Y%m%d-%H%M%S}.docx")

        doc = Document()
        section = doc.sections[0]
        section.top_margin = section.bottom_margin = Inches(0.55)
        section.left_margin = section.right_margin = Inches(0.7)
        styles = doc.styles
        styles["Normal"].font.name = "Aptos"
        styles["Normal"].font.size = Pt(10.5)

        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run(name.upper())
        run.bold = True; run.font.size = Pt(22); run.font.color.rgb = RGBColor(35, 55, 92)
        role = doc.add_paragraph()
        role.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = role.add_run(target_role)
        rr.bold = True; rr.font.size = Pt(12)
        contact_parts = [str(params.get(k) or "").strip() for k in ("email", "phone", "location")]
        if not contact_parts[-1]:
            contact_parts[-1] = defaults["city"]
        contact = "  |  ".join(part for part in contact_parts if part)
        if contact:
            p = doc.add_paragraph(contact); p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        def heading(text: str):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(text.upper()); r.bold = True; r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(53, 83, 145)

        summary = str(params.get("summary") or "").strip()
        if summary:
            heading("Professional Summary"); doc.add_paragraph(summary)

        skills = [str(x).strip() for x in (params.get("skills") or []) if str(x).strip()]
        if skills:
            heading("Skills"); doc.add_paragraph(" • ".join(skills))

        experience = [str(x).strip() for x in (params.get("experience") or []) if str(x).strip()]
        if experience:
            heading("Experience")
            for item in experience:
                doc.add_paragraph(item, style="List Bullet")

        projects = [str(x).strip() for x in (params.get("projects") or defaults["projects"]) if str(x).strip()]
        if projects:
            heading("Projects")
            for item in projects:
                doc.add_paragraph(item, style="List Bullet")

        education = [str(x).strip() for x in (params.get("education") or []) if str(x).strip()]
        if education:
            heading("Education")
            for item in education:
                doc.add_paragraph(item, style="List Bullet")

        doc.save(output)
        push_undo(f"created resume {output.name}", lambda p=output: _remove(p))
        return f"Resume created for {target_role}: {output}"
    except Exception as exc:
        return f"Resume creation failed: {exc}"


TOOL = {
    "name": "resume_builder",
    "description": (
        "Create a polished DOCX resume using remembered identity/project context and a short user brief. "
        "Ask only for a missing full name or target role; optional details can be added later. "
        "Always creates a new file and supports undo."
    ),
    "parameters": {
        "type": "OBJECT",
        "properties": {
            "full_name": {"type": "STRING"}, "target_role": {"type": "STRING"},
            "email": {"type": "STRING"}, "phone": {"type": "STRING"},
            "location": {"type": "STRING"}, "summary": {"type": "STRING"},
            "skills": {"type": "ARRAY", "items": {"type": "STRING"}},
            "experience": {"type": "ARRAY", "items": {"type": "STRING"}},
            "projects": {"type": "ARRAY", "items": {"type": "STRING"}},
            "education": {"type": "ARRAY", "items": {"type": "STRING"}},
            "output_path": {"type": "STRING"},
        },
    },
    "handler": resume_builder,
}
