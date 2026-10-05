"""Precise Excel editing with copy-first safety and undo support."""
from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from core.undo import push_undo


def _unique_output(source: Path) -> Path:
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    return source.with_name(f"{source.stem}_charlie_{stamp}{source.suffix}")


def _remove_created(path: Path) -> str:
    if path.exists():
        path.unlink()
        return f"Removed {path.name}."
    return f"{path.name} was already missing."


def _load(path: Path, data_only: bool = False):
    from openpyxl import load_workbook
    return load_workbook(path, data_only=data_only, keep_vba=path.suffix.lower() == ".xlsm")


def _inspect(path: Path, sheet_name: str = "") -> str:
    wb = _load(path, data_only=False)
    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
    rows = []
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 8),
                            max_col=min(ws.max_column, 12), values_only=True):
        rows.append(" | ".join("" if value is None else str(value) for value in row))
    preview = "\n".join(rows) or "(empty sheet)"
    return (
        f"Workbook: {path.name}\nSheets: {', '.join(wb.sheetnames)}\n"
        f"Selected sheet: {ws.title}; used range: {ws.max_row} rows x {ws.max_column} columns\n"
        f"Preview:\n{preview}"
    )


def excel_worker(parameters: dict, player=None, **_unused) -> str:
    params = parameters or {}
    action = str(params.get("action") or "inspect").strip().lower()
    raw_path = str(params.get("file_path") or "").strip().strip('"')
    path = Path(raw_path).expanduser() if raw_path else None
    if action != "create" and (path is None or not path.is_file()):
        return "Use file_catalog first and provide an existing Excel file_path."
    if path is not None and path.suffix.lower() not in {".xlsx", ".xlsm"}:
        return "Excel editing supports .xlsx and .xlsm files. Convert older .xls files first."

    try:
        if action == "inspect":
            return _inspect(path, str(params.get("sheet") or ""))

        if action == "create":
            from openpyxl import Workbook
            destination = Path(str(params.get("file_path") or "Charlie_Workbook.xlsx")).expanduser()
            if destination.suffix.lower() != ".xlsx":
                destination = destination.with_suffix(".xlsx")
            if destination.exists():
                destination = _unique_output(destination)
            destination.parent.mkdir(parents=True, exist_ok=True)
            wb = Workbook()
            wb.active.title = str(params.get("sheet") or "Sheet1")[:31]
            wb.save(destination)
            push_undo(f"created Excel workbook {destination.name}",
                      lambda p=destination: _remove_created(p))
            return f"Created Excel workbook: {destination}"

        wb = _load(path)
        sheet = str(params.get("sheet") or "").strip()
        ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.active

        if action == "set_cells":
            edits = params.get("edits") or []
            if not isinstance(edits, list) or not edits:
                return "Provide edits as a list of cell/value objects."
            for edit in edits[:500]:
                if not isinstance(edit, dict) or not edit.get("cell"):
                    continue
                ws[str(edit["cell"]).upper()] = edit.get("value")

        elif action == "append_rows":
            rows = params.get("rows") or []
            if not isinstance(rows, list) or not rows:
                return "Provide rows as a list of lists."
            for row in rows[:5000]:
                ws.append(list(row) if isinstance(row, (list, tuple)) else [row])

        elif action == "format":
            from openpyxl.styles import Font, PatternFill, Alignment
            bold_header = bool(params.get("bold_header", True))
            if bold_header and ws.max_row:
                for cell in ws[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="3B5CCC")
                    cell.alignment = Alignment(horizontal="center")
            if bool(params.get("freeze_header", True)):
                ws.freeze_panes = "A2"
            if bool(params.get("auto_width", True)):
                for column in ws.columns:
                    letter = column[0].column_letter
                    width = min(55, max(10, max(len(str(c.value or "")) for c in column) + 2))
                    ws.column_dimensions[letter].width = width
        else:
            return "action must be inspect, create, set_cells, append_rows, or format."

        output = _unique_output(path)
        wb.save(output)
        push_undo(f"created edited Excel copy {output.name}",
                  lambda p=output: _remove_created(p))
        return f"Excel task completed on sheet '{ws.title}'. Saved edited copy: {output}"
    except PermissionError:
        return "Excel could not save the file because it is open or access is denied. Close the workbook and retry."
    except Exception as exc:
        return f"Excel task failed: {exc}"


TOOL = {
    "name": "excel_worker",
    "description": (
        "Inspect, create, edit, append to, and format Excel .xlsx/.xlsm workbooks. "
        "Use file_catalog first when the user did not provide an exact path. "
        "Inspect the workbook before editing. It always saves a new copy and supports undo, "
        "so ordinary edits do not need confirmation. Ask only when the target file, sheet, or "
        "requested transformation is genuinely ambiguous."
    ),
    "parameters": {
        "type": "OBJECT",
        "properties": {
            "action": {"type": "STRING", "description": "inspect | create | set_cells | append_rows | format"},
            "file_path": {"type": "STRING", "description": "Exact workbook path or destination for create"},
            "sheet": {"type": "STRING", "description": "Worksheet name; blank selects the active sheet"},
            "edits": {
                "type": "ARRAY", "description": "Cell changes for set_cells",
                "items": {"type": "OBJECT", "properties": {
                    "cell": {"type": "STRING"}, "value": {"type": "STRING"}
                }, "required": ["cell", "value"]},
            },
            "rows": {"type": "ARRAY", "description": "Rows to append", "items": {"type": "ARRAY", "items": {"type": "STRING"}}},
            "bold_header": {"type": "BOOLEAN"},
            "freeze_header": {"type": "BOOLEAN"},
            "auto_width": {"type": "BOOLEAN"},
        },
        "required": ["action"],
    },
    "handler": excel_worker,
}
